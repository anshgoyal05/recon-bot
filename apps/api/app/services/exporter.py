from pathlib import Path
from io import BytesIO
from typing import List, Dict, Any
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from sqlmodel import Session, select
from app.models.database import engine
from app.models.batch import Batch, Document, ExtractedDoc
from app.models.match import MatchSet, LineItem, Discrepancy

# Apple / Enterprise Color Palette
HEADER_FILL = PatternFill(start_color="1D1D1F", end_color="1D1D1F", fill_type="solid")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

SUCCESS_FILL = PatternFill(start_color="E8F8EE", end_color="E8F8EE", fill_type="solid")
SUCCESS_FONT = Font(name="Calibri", size=10, bold=True, color="1E7E34")

WARNING_FILL = PatternFill(start_color="FFF8E6", end_color="FFF8E6", fill_type="solid")
WARNING_FONT = Font(name="Calibri", size=10, bold=True, color="D38B02")

DANGER_FILL = PatternFill(start_color="FDECEB", end_color="FDECEB", fill_type="solid")
DANGER_FONT = Font(name="Calibri", size=10, bold=True, color="D93025")

TITLE_FONT = Font(name="Calibri", size=16, bold=True, color="1D1D1F")
SUBTITLE_FONT = Font(name="Calibri", size=11, italic=True, color="6E6E73")
BOLD_FONT = Font(name="Calibri", size=10, bold=True)
REGULAR_FONT = Font(name="Calibri", size=10)

THIN_BORDER = Border(
    left=Side(style='thin', color='DDDDDD'),
    right=Side(style='thin', color='DDDDDD'),
    top=Side(style='thin', color='DDDDDD'),
    bottom=Side(style='thin', color='DDDDDD')
)

def export_batch_excel(batch_id: str) -> BytesIO:
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    with Session(engine) as session:
        batch = session.get(Batch, batch_id)
        if not batch:
            ws = wb.create_sheet(title="Report")
            ws.cell(row=1, column=1, value="Batch not found")
            stream = BytesIO()
            wb.save(stream)
            stream.seek(0)
            return stream

        match_sets = session.exec(select(MatchSet).where(MatchSet.batch_id == batch_id)).all()

        # =====================================================================
        # 1. Summary Sheet
        # =====================================================================
        ws_sum = wb.create_sheet(title="Reconciliation Summary")
        ws_sum.views.sheetView[0].showGridLines = True

        # Title Banner
        ws_sum.cell(row=2, column=2, value="Recon-Bot Three-Way Matching Audit Report").font = TITLE_FONT
        ws_sum.cell(row=3, column=2, value=f"Batch: {batch.name} | Generated: {batch.updated_at.strftime('%Y-%m-%d %H:%M:%S UTC')}").font = SUBTITLE_FONT

        # KPI Summary Cards
        kpis = [
            ("Total Invoices Analyzed", len(match_sets)),
            ("Fully Matched (Green)", batch.matched_count),
            ("Requires Review (Amber)", batch.review_count),
            ("Critical Mismatches (Red)", batch.mismatch_count),
            ("Total Rupee-at-Risk (INR)", f"₹{batch.total_at_risk_inr:,.2f}"),
        ]
        
        for k_idx, (label, val) in enumerate(kpis, start=2):
            cell_lbl = ws_sum.cell(row=5, column=k_idx, value=label)
            cell_lbl.font = Font(name="Calibri", size=9, color="6E6E73")
            cell_lbl.alignment = Alignment(horizontal="center")
            
            cell_val = ws_sum.cell(row=6, column=k_idx, value=val)
            cell_val.font = Font(name="Calibri", size=14, bold=True, color="1D1D1F")
            cell_val.alignment = Alignment(horizontal="center")

            # Border for KPI card
            for r in (5, 6):
                ws_sum.cell(row=r, column=k_idx).border = THIN_BORDER
                ws_sum.cell(row=r, column=k_idx).fill = PatternFill(start_color="F5F5F7", end_color="F5F5F7", fill_type="solid")

        # Summary Table Headers
        headers = [
            "Set ID", "Vendor / Supplier", "PO Reference", "Invoice Number",
            "GRN Number", "Audit Verdict", "₹ at Risk", "Discrepancies Count"
        ]
        start_row = 9
        for col_idx, h in enumerate(headers, start=2):
            c = ws_sum.cell(row=start_row, column=col_idx, value=h)
            c.font = HEADER_FONT
            c.fill = HEADER_FILL
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = THIN_BORDER

        # Summary Rows
        for s_idx, mset in enumerate(match_sets, start=start_row + 1):
            discrepancies = session.exec(select(Discrepancy).where(Discrepancy.match_set_id == mset.id)).all()

            ws_sum.cell(row=s_idx, column=2, value=mset.id[:8]).font = REGULAR_FONT
            ws_sum.cell(row=s_idx, column=3, value=mset.supplier_name or "-").font = REGULAR_FONT
            ws_sum.cell(row=s_idx, column=4, value=mset.po_number or "-").font = REGULAR_FONT
            ws_sum.cell(row=s_idx, column=5, value=mset.invoice_number or "-").font = REGULAR_FONT
            ws_sum.cell(row=s_idx, column=6, value=mset.grn_number or "-").font = REGULAR_FONT
            
            v_cell = ws_sum.cell(row=s_idx, column=7, value=mset.verdict)
            if mset.verdict == "MATCHED":
                v_cell.fill = SUCCESS_FILL
                v_cell.font = SUCCESS_FONT
            elif mset.verdict == "REVIEW":
                v_cell.fill = WARNING_FILL
                v_cell.font = WARNING_FONT
            else:
                v_cell.fill = DANGER_FILL
                v_cell.font = DANGER_FONT
            v_cell.alignment = Alignment(horizontal="center")

            risk_cell = ws_sum.cell(row=s_idx, column=8, value=mset.total_at_risk_inr)
            risk_cell.number_format = '₹#,##0.00'
            risk_cell.font = BOLD_FONT
            risk_cell.alignment = Alignment(horizontal="right")

            cnt_cell = ws_sum.cell(row=s_idx, column=9, value=len(discrepancies))
            cnt_cell.font = REGULAR_FONT
            cnt_cell.alignment = Alignment(horizontal="center")

            for c_i in range(2, 10):
                ws_sum.cell(row=s_idx, column=c_i).border = THIN_BORDER

        # Auto-adjust column widths
        for col in ws_sum.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws_sum.column_dimensions[col_letter].width = max(max_len + 3, 12)

        # =====================================================================
        # 2. Per-Set Detail Sheets
        # =====================================================================
        for idx, mset in enumerate(match_sets, start=1):
            sheet_title = f"Set {idx} - {mset.invoice_number or 'Details'}"[:31]
            ws_set = wb.create_sheet(title=sheet_title)
            ws_set.views.sheetView[0].showGridLines = True

            ws_set.cell(row=2, column=2, value=f"{mset.set_name}").font = TITLE_FONT
            ws_set.cell(row=3, column=2, value=f"Supplier: {mset.supplier_name} ({mset.supplier_gstin}) | Verdict: {mset.verdict} | Rupee-at-Risk: ₹{mset.total_at_risk_inr:,.2f}").font = SUBTITLE_FONT

            # Discrepancies Table
            discrepancies = session.exec(select(Discrepancy).where(Discrepancy.match_set_id == mset.id)).all()
            ws_set.cell(row=5, column=2, value="Audit Findings & Discrepancies:").font = BOLD_FONT

            disc_headers = ["Severity", "Category", "Field", "Expected Value", "Billed / Actual", "₹ Impact", "Explanation"]
            for d_c, dh in enumerate(disc_headers, start=2):
                c = ws_set.cell(row=6, column=d_c, value=dh)
                c.fill = HEADER_FILL
                c.font = HEADER_FONT
                c.border = THIN_BORDER

            d_row = 7
            if not discrepancies:
                ws_set.cell(row=d_row, column=2, value="No discrepancies found. Three-way match verified 100% clean.").font = SUCCESS_FONT
                ws_set.merge_cells(start_row=d_row, start_column=2, end_row=d_row, end_column=8)
                d_row += 1
            else:
                for d in discrepancies:
                    sev_c = ws_set.cell(row=d_row, column=2, value=d.severity)
                    if d.severity == "ERROR":
                        sev_c.fill = DANGER_FILL
                        sev_c.font = DANGER_FONT
                    else:
                        sev_c.fill = WARNING_FILL
                        sev_c.font = WARNING_FONT
                    sev_c.alignment = Alignment(horizontal="center")

                    ws_set.cell(row=d_row, column=3, value=d.discrepancy_type).font = REGULAR_FONT
                    ws_set.cell(row=d_row, column=4, value=d.field_name).font = REGULAR_FONT
                    ws_set.cell(row=d_row, column=5, value=d.expected_value).font = REGULAR_FONT
                    ws_set.cell(row=d_row, column=6, value=d.actual_value).font = REGULAR_FONT
                    
                    imp_c = ws_set.cell(row=d_row, column=7, value=d.impact_inr)
                    imp_c.number_format = '₹#,##0.00'
                    imp_c.font = BOLD_FONT
                    
                    ws_set.cell(row=d_row, column=8, value=d.explanation).font = REGULAR_FONT

                    for ci in range(2, 9):
                        ws_set.cell(row=d_row, column=ci).border = THIN_BORDER
                    d_row += 1

            for col in ws_set.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = get_column_letter(col[0].column)
                ws_set.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream
